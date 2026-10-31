#!/usr/bin/env python3
"""프론트 레이아웃 인벤토리에서 사이트 구조도 산출물을 만든다.

입력: 실제 화면을 확인해 정리한 정규화 JSON
출력: site-structure.xlsx, site-structure.md, site-structure.html

API 목록은 화면 구조의 근거가 아니다. 페이지 계층은 menu_path와 parent_id로
표현하고, API는 화면별 매핑과 API only 부록으로만 기록한다.
"""

from __future__ import annotations

import argparse
import html
import json
import re
import shlex
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

OUTPUT_FORMATS = ("md", "xlsx", "html")
DEFAULT_FORMATS = "xlsx,md,html"
STATUSES = {"active", "planned", "deprecated"}
REVIEW_QUESTIONS = [
    "메뉴 depth가 실제 프론트 메뉴 구조와 일치하는가?",
    "상세·탭·액션 화면이 부모 화면으로 연결됐는가?",
    "API only가 화면 구조도와 분리됐는가?",
    "스크린샷 경로가 실제 화면과 일치하는가?",
    "KRDS 메인 메뉴 최대 3Depth 기준을 지켰는가?",
]


def parse_args():
    parser = argparse.ArgumentParser(description="사이트 구조도 산출물 생성")
    parser.add_argument("--input", required=True, help="정규화된 사이트 구조 JSON")
    parser.add_argument("--outdir", default="docs/site-structure", help="출력 디렉터리")
    parser.add_argument("--name", default="site-structure", help="출력 파일 기본 이름")
    parser.add_argument(
        "--formats",
        default=DEFAULT_FORMATS,
        help=f"생성할 산출물: {','.join(OUTPUT_FORMATS)} (기본: {DEFAULT_FORMATS})",
    )
    parser.add_argument(
        "--sources",
        default="",
        help="산출물 갱신 판단에 쓸 소스 글롭 (쉼표 구분). 기본은 입력 JSON과 src/app·src/pages·src/routes",
    )
    parser.add_argument(
        "--no-manifest",
        action="store_true",
        help="sidecar .artifact.json manifest를 만들지 않음",
    )
    parser.add_argument(
        "--strict",
        action="store_true",
        help="검증 경고가 있으면 종료 코드 1",
    )
    args = parser.parse_args()
    formats = [item.strip().lower() for item in args.formats.split(",") if item.strip()]
    unknown = [item for item in formats if item not in OUTPUT_FORMATS]
    if unknown:
        parser.error(f"지원하지 않는 형식입니다: {', '.join(unknown)}")
    if not formats:
        parser.error("--formats 에 하나 이상의 형식을 지정하세요.")
    args.formats = set(formats)
    return args


def repo_path(path):
    candidate = Path(path)
    if not candidate.is_absolute():
        candidate = (Path.cwd() / candidate).resolve()
    else:
        candidate = candidate.resolve()
    try:
        return candidate.relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return candidate.as_posix()


def git_head_commit():
    try:
        result = subprocess.run(
            ["git", "rev-parse", "--short", "HEAD"],
            capture_output=True,
            text=True,
            check=False,
        )
    except OSError:
        return ""
    return result.stdout.strip() if result.returncode == 0 else ""


def git_dirty():
    try:
        result = subprocess.run(
            ["git", "status", "--porcelain"],
            capture_output=True,
            text=True,
            check=False,
        )
    except OSError:
        return False
    return bool(result.stdout.strip()) if result.returncode == 0 else False


PATH_OPTIONS = {"--input", "--outdir"}


def portable_value(value):
    if not value or value == "-" or value.startswith("-"):
        return value
    candidate = Path(value)
    if not candidate.is_absolute():
        return value
    try:
        return candidate.expanduser().resolve().relative_to(Path.cwd().resolve()).as_posix()
    except (ValueError, OSError):
        return value


def portable_args(argv):
    result = []
    index = 0
    while index < len(argv):
        item = argv[index]
        if item in PATH_OPTIONS and index + 1 < len(argv):
            result.extend([item, portable_value(argv[index + 1])])
            index += 2
            continue
        matched = False
        for option in PATH_OPTIONS:
            if item.startswith(option + "="):
                result.append(option + "=" + portable_value(item[len(option) + 1 :]))
                matched = True
                break
        if not matched:
            result.append(item)
        index += 1
    return result


def regenerate_command():
    script = Path(sys.argv[0])
    if script.exists():
        relative = repo_path(script)
        if not Path(relative).is_absolute() and not relative.startswith("../"):
            return "python3 " + shlex.join([relative, *portable_args(sys.argv[1:])])
    if script.parent.name == "scripts" and script.parent.parent.name:
        portable = f"{script.parent.parent.name}/scripts/{script.name}"
    else:
        portable = script.name
    return "python3 " + shlex.join([portable, *portable_args(sys.argv[1:])])


def write_manifest(path, artifact, outputs, sources):
    payload = {
        "version": 1,
        "skill": "gorani-site-structure",
        "artifact": artifact,
        "outputs": outputs,
        "sources": sources,
        "generated_commit": git_head_commit(),
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "working_tree_dirty": git_dirty(),
        "regenerate": regenerate_command(),
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def as_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    return [text] if text else []


def normalize_menu_path(value, title):
    if isinstance(value, list):
        parts = [str(item).strip() for item in value if str(item).strip()]
    else:
        text = str(value or "").strip()
        parts = [item.strip() for item in re.split(r"\s*>\s*", text) if item.strip()]
    return parts or [title]


def normalize_roles(value):
    if isinstance(value, list):
        values = [str(item).strip() for item in value if str(item).strip()]
    else:
        values = [item.strip() for item in str(value or "").split(",") if item.strip()]
    return values or ["전체"]


def normalize_bool(value, default=True):
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    text = str(value).strip().lower()
    if text in {"y", "yes", "true", "1", "예"}:
        return True
    if text in {"n", "no", "false", "0", "아니오"}:
        return False
    return default


def normalize_api(raw):
    if isinstance(raw, str):
        text = raw.strip()
        if not text:
            return None
        parts = text.split(None, 1)
        method = parts[0].upper()
        path = parts[1] if len(parts) > 1 else ""
        return {"method": method, "path": path, "purpose": "", "when": "", "note": ""}
    if not isinstance(raw, dict):
        return None
    method = str(raw.get("method", "")).strip().upper()
    path = str(raw.get("path", "")).strip()
    if not method or not path:
        return None
    return {
        "method": method,
        "path": path,
        "purpose": str(raw.get("purpose", "")).strip(),
        "when": str(raw.get("when", "")).strip(),
        "note": str(raw.get("note", "")).strip(),
    }


def load_input(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    raw_pages = data.get("pages")
    if not isinstance(raw_pages, list) or not raw_pages:
        sys.exit("입력 JSON의 pages 배열이 비어 있습니다.")
    pages = []
    ids = set()
    for raw in raw_pages:
        if not isinstance(raw, dict):
            continue
        page_id = str(raw.get("id", "")).strip()
        if not page_id:
            sys.exit("page.id 는 필수입니다.")
        if page_id in ids:
            sys.exit(f"page.id 가 중복됐습니다: {page_id}")
        ids.add(page_id)
        title = str(raw.get("title", "")).strip() or page_id
        status = str(raw.get("status", "active")).strip().lower() or "active"
        if status not in STATUSES:
            sys.exit(f"page.status 는 active/planned/deprecated 중 하나여야 합니다: {page_id}")
        apis = [api for api in (normalize_api(item) for item in raw.get("apis", []) or []) if api]
        pages.append(
            {
                "id": page_id,
                "title": title,
                "menu_path": normalize_menu_path(raw.get("menu_path"), title),
                "parent_id": str(raw.get("parent_id", "")).strip(),
                "url": str(raw.get("url", "")).strip(),
                "layout": str(raw.get("layout", "")).strip(),
                "roles": normalize_roles(raw.get("roles")),
                "exposure": str(raw.get("exposure", "")).strip() or "미확인",
                "desktop": normalize_bool(raw.get("desktop"), True),
                "mobile": normalize_bool(raw.get("mobile"), True),
                "status": status,
                "screenshot": str(raw.get("screenshot", "")).strip(),
                "apis": apis,
                "source": str(raw.get("source", "")).strip(),
                "note": str(raw.get("note", "")).strip(),
            }
        )
    if not pages:
        sys.exit("정규화된 page가 없습니다.")

    page_map = {page["id"]: page for page in pages}
    warnings = []
    for raw in data.get("warnings", []) or []:
        if isinstance(raw, dict) and str(raw.get("message", "")).strip():
            warnings.append(
                {
                    "category": str(raw.get("category", "검토")).strip() or "검토",
                    "message": str(raw.get("message", "")).strip(),
                    "action": str(raw.get("action", "")).strip(),
                }
            )
    depth_count = max(2, max(len(page["menu_path"]) for page in pages))
    if depth_count > 3:
        warnings.append(
            {
                "category": "KRDS",
                "message": f"메뉴 depth가 {depth_count}수준입니다. KRDS 메인 메뉴는 최대 3수준을 권장합니다.",
                "action": "3Depth 초과 항목을 메뉴가 아닌 상세·탭·액션 화면으로 분리하세요.",
            }
        )
    for page in pages:
        page["parent_title"] = ""
        if page["parent_id"]:
            parent = page_map.get(page["parent_id"])
            if parent is None:
                warnings.append(
                    {
                        "category": "구조",
                        "message": f"{page['id']}의 parent_id가 없습니다: {page['parent_id']}",
                        "action": "부모 화면 ID를 확인하세요.",
                    }
                )
            else:
                page["parent_title"] = parent["title"]
        if page["status"] == "active" and not page["screenshot"]:
            warnings.append(
                {
                    "category": "증적",
                    "message": f"{page['id']}의 스크린샷 경로가 없습니다.",
                    "action": "실제 화면을 캡처해 screens/ 아래에 저장하세요.",
                }
            )
    standards = []
    for raw in data.get("standards", []) or []:
        if not isinstance(raw, dict):
            continue
        name = str(raw.get("name", "")).strip()
        if not name:
            continue
        standards.append(
            {
                "source": str(raw.get("source", "")).strip(),
                "name": name,
                "scope": str(raw.get("scope", "")).strip(),
                "rule": str(raw.get("rule", "")).strip(),
                "url": str(raw.get("url", "")).strip(),
            }
        )
    api_only = []
    for raw in data.get("api_only", []) or []:
        api = normalize_api(raw)
        if api:
            api["caller"] = str(raw.get("caller", "")).strip()
            api["reason"] = str(raw.get("reason", "")).strip()
            api_only.append(api)

    meta = {
        "site": str(data.get("site", "")).strip() or "Site",
        "base_url": str(data.get("base_url", "")).strip(),
        "inspected_at": str(data.get("inspected_at", "")).strip(),
        "source": str(data.get("source", "")).strip(),
        "standard": str(data.get("standard", "")).strip(),
        "roles": as_list(data.get("roles")),
        "viewports": as_list(data.get("viewports")),
        "review_status": str(data.get("review_status", "")).strip(),
        "sources": as_list(data.get("sources")),
    }
    return meta, pages, depth_count, warnings, standards, api_only


def table_label(page):
    return " > ".join(page["menu_path"])


def status_counts(pages):
    return Counter(page["status"] for page in pages)


def api_count(pages):
    return sum(len(page["apis"]) for page in pages)


def md_cell(value):
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ") or "-"


def bool_label(value):
    return "Y" if value else "N"


def build_menu_tree(pages):
    root = {"label": "", "children": {}, "pages": []}
    for page in pages:
        node = root
        for label in page["menu_path"]:
            node = node["children"].setdefault(label, {"label": label, "children": {}, "pages": []})
        node["pages"].append(page)
    return root


def build_markdown(meta, pages, depth_count, warnings, standards, api_only):
    counts = status_counts(pages)
    headers = ["No"] + [f"{index}Depth" for index in range(1, depth_count + 1)]
    headers += ["화면 ID", "화면명", "부모 화면", "노출", "URL 패턴", "레이아웃", "접근 역할", "데스크톱", "모바일", "상태", "스크린샷 경로"]
    lines = [
        "# 사이트 구조도",
        "",
        f"- 사이트: {meta['site']}",
        f"- 기준 URL: {meta['base_url'] or '-'}",
        f"- 검사 시점: {meta['inspected_at'] or '-'}",
        f"- 수집 기준: {meta['source'] or '-'}",
        f"- 적용 기준: {meta['standard'] or '-'}",
        f"- 검사 역할: {', '.join(meta['roles']) or '-'}",
        f"- 검사 화면 폭: {', '.join(meta['viewports']) or '-'}",
        f"- Depth 규칙: 최소 2Depth, 최대 3Depth (KRDS 메인 메뉴 기준), 현재 {depth_count}Depth",
        f"- 화면: {len(pages)}",
        f"- 활성/계획/폐기: {counts.get('active', 0)} / {counts.get('planned', 0)} / {counts.get('deprecated', 0)}",
        f"- API 매핑: {api_count(pages)}",
        f"- API only: {len(api_only)}",
        f"- 검토 메모: {len(warnings)}",
        "",
        "> 이 문서는 실제 프론트 레이아웃을 확인해 정리한 구조도입니다. API 목록은 화면 구조의 근거가 아니며, 화면별 연결과 API only 부록으로만 사용합니다.",
        "",
        "## 사이트 구조",
        "",
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for index, page in enumerate(pages, start=1):
        depths = page["menu_path"] + [""] * (depth_count - len(page["menu_path"]))
        row = [str(index)] + [md_cell(value or "-") for value in depths]
        row += [
            md_cell(page["id"]),
            md_cell(page["title"]),
            md_cell(page["parent_title"] or "-"),
            md_cell(page["exposure"]),
            md_cell(page["url"] or "-"),
            md_cell(page["layout"] or "-"),
            md_cell(", ".join(page["roles"])),
            bool_label(page["desktop"]),
            bool_label(page["mobile"]),
            md_cell(page["status"]),
            md_cell(page["screenshot"] or "-"),
        ]
        lines.append("| " + " | ".join(row) + " |")
    lines += [
        "",
        "## 화면 목록",
        "",
        "| 화면 ID | 화면명 | 메뉴 경로 | URL 패턴 | 부모 화면 | 레이아웃 | 접근 역할 | 노출 방식 | 데스크톱 | 모바일 | 상태 | 스크린샷 경로 | API 호출 | 비고 |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for page in pages:
        apis = "<br>".join(f"{api['method']} {api['path']}" for api in page["apis"]) or "-"
        lines.append(
            "| "
            + " | ".join(
                [
                    md_cell(page["id"]),
                    md_cell(page["title"]),
                    md_cell(table_label(page)),
                    md_cell(page["url"] or "-"),
                    md_cell(page["parent_title"] or "-"),
                    md_cell(page["layout"] or "-"),
                    md_cell(", ".join(page["roles"])),
                    md_cell(page["exposure"]),
                    bool_label(page["desktop"]),
                    bool_label(page["mobile"]),
                    md_cell(page["status"]),
                    md_cell(page["screenshot"] or "-"),
                    md_cell(apis),
                    md_cell(page["note"] or "-"),
                ]
            )
            + " |"
        )
    lines += ["", "## API 매핑", ""]
    if api_count(pages):
        lines += ["| 화면 ID | 화면명 | Method | API 경로 | 호출 시점 | 목적 |", "| --- | --- | --- | --- | --- | --- |"]
        for page in pages:
            for api in page["apis"]:
                when = api["when"] or ("on load" if api["method"] == "GET" else "사용자 액션")
                lines.append(
                    f"| {md_cell(page['id'])} | {md_cell(page['title'])} | {md_cell(api['method'])} | {md_cell(api['path'])} | {md_cell(when)} | {md_cell(api['purpose'] or '-')} |"
                )
    else:
        lines.append("- API 매핑 없음")
    lines += ["", "## API only", ""]
    if api_only:
        lines += ["| Method | API 경로 | 용도 | 호출 주체 | 화면 없음 사유 |", "| --- | --- | --- | --- | --- |"]
        for api in api_only:
            lines.append(
                f"| {md_cell(api['method'])} | {md_cell(api['path'])} | {md_cell(api['purpose'] or '-')} | {md_cell(api['caller'] or '-')} | {md_cell(api['reason'] or '-')} |"
            )
    else:
        lines.append("- API only 없음")
    lines += ["", "## 표준 근거", ""]
    if standards:
        lines += ["| 구분 | 문서 | 적용 항목 | 내용 | 링크 |", "| --- | --- | --- | --- | --- |"]
        for standard in standards:
            link = f"[{md_cell(standard['name'])}]({standard['url']})" if standard["url"] else md_cell(standard["name"])
            lines.append(
                f"| {md_cell(standard['source'] or '-')} | {link} | {md_cell(standard['scope'] or '-')} | {md_cell(standard['rule'] or '-')} | {md_cell(standard['url'] or '-')} |"
            )
    else:
        lines.append("- 표준 근거 없음")
    lines += ["", "## 검토 메모", ""]
    if warnings:
        lines += ["| No | 구분 | 내용 | 확인 필요 |", "| --- | --- | --- | --- |"]
        for index, warning in enumerate(warnings, start=1):
            lines.append(
                f"| {index} | {md_cell(warning['category'])} | {md_cell(warning['message'])} | {md_cell(warning['action'] or '-')} |"
            )
    else:
        lines.append("- 검토 메모 없음")
    lines += ["", "## 검토 요청", ""]
    for index, question in enumerate(REVIEW_QUESTIONS, start=1):
        lines.append(f"{index}. {question}")
    return "\n".join(lines).rstrip() + "\n"


def write_xlsx(path, meta, pages, depth_count, warnings, standards, api_only):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit(
            "xlsx를 만들려면 openpyxl이 필요합니다. 스킬의 도구 설치 원칙에 따라 설치한 뒤 다시 시도하세요."
        )

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(size=16, bold=True)
    section_font = Font(bold=True)
    group_fills = [
        PatternFill("solid", fgColor="D9E1F2"),
        PatternFill("solid", fgColor="E2EFDA"),
        PatternFill("solid", fgColor="FFF2CC"),
        PatternFill("solid", fgColor="E4DFEC"),
        PatternFill("solid", fgColor="FCE4D6"),
        PatternFill("solid", fgColor="DDEBF7"),
    ]
    active_fill = PatternFill("solid", fgColor="C6E0B4")
    planned_fill = PatternFill("solid", fgColor="FFE699")
    deprecated_fill = PatternFill("solid", fgColor="D9D9D9")
    api_fill = PatternFill("solid", fgColor="E4DFEC")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(ws, widths, columns):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A2"

    def append_row(ws, values):
        ws.append(values)
        row = ws.max_row
        for cell in ws[row]:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "요약"
    summary.sheet_properties.tabColor = "1F3864"
    summary["A1"] = f"사이트 구조도 — {meta['site']}"
    summary["A1"].font = title_font
    summary.merge_cells("A1:D1")
    summary["A3"] = "기본 정보"
    summary["A3"].font = section_font
    counts = status_counts(pages)
    rows = [
        ("사이트", meta["site"]),
        ("기준 URL", meta["base_url"] or "-"),
        ("검사 시점", meta["inspected_at"] or "-"),
        ("수집 기준", meta["source"] or "-"),
        ("적용 기준", meta["standard"] or "-"),
        ("검사 역할", ", ".join(meta["roles"]) or "-"),
        ("검사 화면 폭", ", ".join(meta["viewports"]) or "-"),
        ("Depth 규칙", f"메뉴 성격의 depth만 사용. 최소 2Depth, 최대 3Depth (KRDS 메인 메뉴 기준). 현재 {depth_count}Depth"),
        ("수집 방식", "에이전트 화면 확인 + 선택적 Playwright"),
        ("계층 기준", "메뉴 계층 우선, 3Depth 초과는 상세·탭·액션으로 분리"),
        ("화면 수", len(pages)),
        ("활성/계획/폐기", f"{counts.get('active', 0)} / {counts.get('planned', 0)} / {counts.get('deprecated', 0)}"),
        ("내비게이션 그룹", len({page["menu_path"][0] for page in pages})),
        ("API 매핑", api_count(pages)),
        ("API only", len(api_only)),
        ("표준 근거", len(standards)),
        ("검토 메모", len(warnings)),
        ("검토 상태", meta["review_status"] or "초안 — 검토 필요"),
    ]
    for index, (label, value) in enumerate(rows, start=4):
        summary.cell(row=index, column=1, value=label).font = section_font
        summary.cell(row=index, column=2, value=value)
    start = len(rows) + 6
    summary.cell(row=start, column=1, value="시트 구성").font = section_font
    sheet_notes = [
        ("사이트 구조", "실제 메뉴 계층을 Depth 컬럼으로 표시"),
        ("화면 목록", "화면별 URL·역할·노출·상태·스크린샷·API"),
        ("API 매핑", "화면과 API 호출의 연결"),
        ("API only", "화면 없이 API만 있는 경로. 구조도와 분리"),
        ("표준 근거", "KRDS·전자정부 품질관리 지침의 적용 항목"),
        ("검토 메모", "확인 필요 사항과 가정"),
    ]
    for index, (name, desc) in enumerate(sheet_notes, start=start + 1):
        summary.cell(row=index, column=1, value=name).font = section_font
        summary.cell(row=index, column=2, value=desc)
    review_row = start + len(sheet_notes) + 3
    summary.cell(row=review_row, column=1, value="검토 요청").font = section_font
    for index, question in enumerate(REVIEW_QUESTIONS, start=1):
        row = review_row + index
        summary.cell(row=row, column=1, value=f"{index}. {question}")
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 70
    summary.column_dimensions["C"].width = 22
    summary.column_dimensions["D"].width = 22
    summary.freeze_panes = "A4"

    ws = workbook.create_sheet("사이트 구조")
    ws.sheet_properties.tabColor = "2F5597"
    headers = ["No"] + [f"{index}Depth" for index in range(1, depth_count + 1)]
    headers += ["화면 ID", "화면명", "부모 화면", "노출", "URL 패턴", "레이아웃", "접근 역할", "데스크톱", "모바일", "상태", "스크린샷 경로"]
    ws.append(headers)
    widths = [6] + [20] * depth_count + [20, 20, 20, 14, 32, 14, 20, 10, 10, 10, 32]
    style_header(ws, widths, headers)
    for index, page in enumerate(pages, start=1):
        depths = page["menu_path"] + [""] * (depth_count - len(page["menu_path"]))
        append_row(
            ws,
            [index] + depths + [
                page["id"],
                page["title"],
                page["parent_title"] or "",
                page["exposure"],
                page["url"],
                page["layout"],
                ", ".join(page["roles"]),
                bool_label(page["desktop"]),
                bool_label(page["mobile"]),
                page["status"],
                page["screenshot"],
            ],
        )
        row = ws.max_row
        ws.cell(row=row, column=2).font = Font(bold=True)
        ws.cell(row=row, column=2).fill = group_fills[(index - 1) % len(group_fills)]
        status_cell = ws.cell(row=row, column=2 + depth_count + 9)
        status_cell.fill = active_fill if page["status"] == "active" else planned_fill if page["status"] == "planned" else deprecated_fill
    ws.auto_filter.ref = ws.dimensions

    ws = workbook.create_sheet("화면 목록")
    ws.sheet_properties.tabColor = "548235"
    headers = ["화면 ID", "화면명", "메뉴 경로", "URL 패턴", "부모 화면", "레이아웃", "접근 역할", "노출 방식", "데스크톱", "모바일", "상태", "스크린샷 경로", "API 호출", "비고"]
    ws.append(headers)
    style_header(ws, [20, 20, 38, 32, 20, 14, 20, 14, 10, 10, 10, 32, 46, 30], headers)
    for page in pages:
        apis = "\n".join(f"{api['method']} {api['path']}" for api in page["apis"])
        append_row(
            ws,
            [
                page["id"],
                page["title"],
                table_label(page),
                page["url"],
                page["parent_title"],
                page["layout"],
                ", ".join(page["roles"]),
                page["exposure"],
                bool_label(page["desktop"]),
                bool_label(page["mobile"]),
                page["status"],
                page["screenshot"],
                apis,
                page["note"],
            ],
        )
        row = ws.max_row
        ws.cell(row=row, column=11).fill = active_fill if page["status"] == "active" else planned_fill if page["status"] == "planned" else deprecated_fill
    ws.auto_filter.ref = ws.dimensions

    ws = workbook.create_sheet("API 매핑")
    ws.sheet_properties.tabColor = "7F6000"
    headers = ["화면 ID", "화면명", "Method", "API 경로", "호출 시점", "목적", "비고"]
    ws.append(headers)
    style_header(ws, [20, 20, 10, 42, 16, 42, 30], headers)
    for page in pages:
        for api in page["apis"]:
            when = api["when"] or ("on load" if api["method"] == "GET" else "사용자 액션")
            append_row(ws, [page["id"], page["title"], api["method"], api["path"], when, api["purpose"], api["note"]])
    ws.auto_filter.ref = ws.dimensions

    ws = workbook.create_sheet("API only")
    ws.sheet_properties.tabColor = "833C00"
    headers = ["Method", "API 경로", "용도", "호출 주체", "화면 없음 사유", "비고"]
    ws.append(headers)
    style_header(ws, [10, 40, 36, 24, 40, 20], headers)
    for api in api_only:
        append_row(ws, [api["method"], api["path"], api["purpose"], api["caller"], api["reason"], api["note"]])
        for cell in ws[ws.max_row]:
            cell.fill = api_fill
    ws.auto_filter.ref = ws.dimensions

    ws = workbook.create_sheet("표준 근거")
    ws.sheet_properties.tabColor = "7030A0"
    headers = ["구분", "문서", "적용 항목", "내용", "링크"]
    ws.append(headers)
    style_header(ws, [12, 44, 14, 68, 90], headers)
    for standard in standards:
        append_row(ws, [standard["source"], standard["name"], standard["scope"], standard["rule"], standard["url"]])
        link_cell = ws.cell(row=ws.max_row, column=5)
        if standard["url"]:
            link_cell.hyperlink = standard["url"]
            link_cell.font = Font(color="1F5FBF", underline="single")
    ws.auto_filter.ref = ws.dimensions

    ws = workbook.create_sheet("검토 메모")
    ws.sheet_properties.tabColor = "7F7F7F"
    headers = ["No", "구분", "내용", "확인 필요"]
    ws.append(headers)
    style_header(ws, [6, 12, 78, 42], headers)
    for index, warning in enumerate(warnings, start=1):
        append_row(ws, [index, warning["category"], warning["message"], warning["action"]])
    ws.auto_filter.ref = ws.dimensions

    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
    workbook.properties.title = "사이트 구조도"
    workbook.properties.subject = f"Site structure — {meta['site']}"
    workbook.properties.creator = "gorani"
    workbook.save(path)


def build_html(meta, pages, depth_count, warnings, standards, api_only):
    document_title = f"사이트 구조도 — {meta['site']}"
    review_questions = "".join(f"<li>{html.escape(question)}</li>" for question in REVIEW_QUESTIONS)
    tree = build_menu_tree(pages)

    def badge(value):
        return html.escape(value or "-")

    def render_tree(node):
        parts = []
        for label, child in node["children"].items():
            parts.append(f'<li class="menu-node"><span class="menu-label">{html.escape(label)}</span><ul>')
            parts.extend(render_tree(child))
            parts.append("</ul></li>")
        for page in node["pages"]:
            search = " ".join([page["title"], " ".join(page["menu_path"]), page["url"], page["exposure"]]).lower()
            parts.append(
                f'<li class="screen-item" data-tree-search="{html.escape(search, quote=True)}">'
                f'<a href="#screen-{html.escape(page["id"], quote=True)}">{html.escape(page["title"])}</a>'
                f'<span class="badge">{badge(page["exposure"])}</span>'
                f'<span class="path">{html.escape(page["url"] or "-")}</span></li>'
            )
        return parts

    tree_html = "".join(render_tree(tree))
    css = """
:root { color-scheme: light dark; --bg:#fff; --fg:#17181c; --muted:#5c6270; --line:#e2e5ea; --panel:#f6f7f9; --accent:#1f5fbf; }
@media (prefers-color-scheme: dark) { :root { --bg:#111216; --fg:#e8eaee; --muted:#9aa1ae; --line:#2a2d35; --panel:#181a20; --accent:#9db8e8; } }
* { box-sizing: border-box; }
body { margin:0; background:var(--bg); color:var(--fg); font:16px/1.55 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,"Noto Sans KR",sans-serif; }
main, header, footer { max-width:86rem; margin:0 auto; padding:0 1.25rem; }
header { padding-top:2.5rem; }
h1 { font-size:1.8rem; margin:0 0 .35rem; letter-spacing:-.02em; }
h2 { font-size:1.15rem; margin:0 0 .75rem; }
h3 { font-size:.95rem; margin:1rem 0 .5rem; }
p { margin:.35rem 0; }
.eyebrow { color:var(--accent); font-size:.72rem; font-weight:700; letter-spacing:.1em; text-transform:uppercase; }
.meta { color:var(--muted); font-size:.875rem; }
.stats { display:grid; grid-template-columns:repeat(auto-fit,minmax(9rem,1fr)); gap:.65rem; margin:1.5rem 0 1rem; }
.stat { padding:.75rem .9rem; background:var(--panel); border:1px solid var(--line); border-radius:.6rem; }
.stat strong { display:block; font-size:1.25rem; }
.stat span { color:var(--muted); font-size:.78rem; }
.controls { display:flex; flex-wrap:wrap; align-items:center; gap:.75rem; margin:1.2rem 0 .8rem; padding:.85rem 1rem; background:var(--panel); border:1px solid var(--line); border-radius:.7rem; }
.controls label { font-weight:600; font-size:.875rem; }
.controls input { flex:1 1 16rem; min-width:0; padding:.55rem .75rem; color:var(--fg); background:var(--bg); border:1px solid var(--line); border-radius:.45rem; font:inherit; }
.controls input:focus-visible { outline:2px solid var(--accent); outline-offset:1px; }
.count { color:var(--muted); font-size:.8125rem; }
.tree { margin:0; padding:1rem 1rem 1rem 1.4rem; background:var(--panel); border:1px solid var(--line); border-radius:.6rem; }
.tree ul { margin:.25rem 0 .25rem 1rem; padding-left:1rem; border-left:1px solid var(--line); }
.tree li { margin:.3rem 0; }
.menu-label { font-weight:700; }
.menu-node > .menu-label { color:var(--fg); }
.screen-item a { color:var(--accent); text-decoration:none; }
.screen-item a:hover { text-decoration:underline; }
.badge { display:inline-block; margin-left:.45rem; padding:.02rem .4rem; color:var(--muted); border:1px solid var(--line); border-radius:999px; font-size:.7rem; }
.path { margin-left:.45rem; color:var(--muted); font: .75rem/1.3 ui-monospace,SFMono-Regular,Menlo,Consolas,monospace; }
.section { margin:1.5rem 0; }
.table-wrap { overflow-x:auto; border:1px solid var(--line); border-radius:.5rem; }
table { width:100%; border-collapse:collapse; font-size:.84rem; }
th, td { padding:.5rem .6rem; border-bottom:1px solid var(--line); text-align:left; vertical-align:top; }
tr:last-child td { border-bottom:0; }
th { background:var(--panel); color:var(--muted); font-size:.76rem; font-weight:600; white-space:nowrap; }
.status { display:inline-block; padding:.02rem .4rem; border-radius:999px; font-size:.72rem; }
.status.active { background:#C6E0B4; color:#1e3b12; }
.status.planned { background:#FFE699; color:#4d3b00; }
.status.deprecated { background:#D9D9D9; color:#333; }
.warnings { padding:1rem; border:1px dashed var(--line); border-radius:.6rem; }
.warnings ul { margin:.25rem 0 0; padding-left:1.2rem; color:var(--muted); font-size:.875rem; }
.links a { color:var(--accent); }
footer { padding:2rem 0 3rem; color:var(--muted); font-size:.8125rem; border-top:1px solid var(--line); }
[hidden] { display:none !important; }
@media print { .controls, footer { display:none; } .section { break-inside:avoid-page; } body { font-size:10.5pt; } }
""".strip()
    javascript = """
(function () {
  var input = document.getElementById("search");
  if (!input) return;
  var count = document.getElementById("result-count");
  var rows = Array.prototype.slice.call(document.querySelectorAll("[data-screen-search]"));
  var treeItems = Array.prototype.slice.call(document.querySelectorAll("[data-tree-search]"));
  var menus = Array.prototype.slice.call(document.querySelectorAll(".menu-node")).reverse();
  function apply() {
    var q = (input.value || "").trim().toLowerCase();
    var shown = 0;
    rows.forEach(function (item) {
      var match = !q || (item.getAttribute("data-screen-search") || "").indexOf(q) !== -1;
      item.hidden = !match;
      if (match) shown += 1;
    });
    treeItems.forEach(function (item) {
      var match = !q || (item.getAttribute("data-tree-search") || "").indexOf(q) !== -1;
      item.hidden = !match;
    });
    menus.forEach(function (menu) {
      var has = menu.querySelector("[data-tree-search]:not([hidden])");
      menu.hidden = !has;
    });
    if (count) count.textContent = q ? shown + " matches" : shown + " screens";
  }
  input.addEventListener("input", apply);
  apply();
})();
""".strip()
    parts = [
        "<!doctype html>",
        '<html lang="ko">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        "<meta http-equiv=\"Content-Security-Policy\" content=\"default-src 'none'; style-src 'unsafe-inline'; script-src 'unsafe-inline'; base-uri 'none'\">",
        f"<title>{html.escape(document_title)}</title>",
        f"<style>{css}</style>",
        "</head>",
        "<body>",
        "<header>",
        '<p class="eyebrow">Information architecture</p>',
        f"<h1>{html.escape(document_title)}</h1>",
        f'<p class="meta">{html.escape(meta["site"])} · {html.escape(meta["inspected_at"] or "-")} · {html.escape(meta["source"] or "-")}</p>',
        f'<p class="meta">역할: {html.escape(", ".join(meta["roles"]) or "-")} · 화면 폭: {html.escape(", ".join(meta["viewports"]) or "-")}</p>',
        '<div class="stats">',
        f'<div class="stat"><strong>{len(pages)}</strong><span>화면</span></div>',
        f'<div class="stat"><strong>{len({page["menu_path"][0] for page in pages})}</strong><span>메뉴 그룹</span></div>',
        f'<div class="stat"><strong>{api_count(pages)}</strong><span>API 매핑</span></div>',
        f'<div class="stat"><strong>{len(api_only)}</strong><span>API only</span></div>',
        "</div>",
        "</header>",
        "<main>",
        f'<div class="controls"><label for="search">화면 검색</label><input id="search" type="search" placeholder="메뉴·화면·URL" autocomplete="off"><p class="count" id="result-count">{len(pages)} screens</p></div>',
        '<section class="section"><h2>메뉴 구조</h2><ul class="tree">',
        tree_html,
        "</ul></section>",
        '<section class="section"><h2>화면 목록</h2><div class="table-wrap"><table>',
        "<thead><tr><th>화면</th><th>메뉴 경로</th><th>URL</th><th>부모 화면</th><th>레이아웃</th><th>노출</th><th>역할</th><th>데스크톱</th><th>모바일</th><th>상태</th><th>스크린샷 경로</th></tr></thead><tbody>",
    ]
    for page in pages:
        search = " ".join([page["title"], " ".join(page["menu_path"]), page["url"], page["exposure"]]).lower()
        status_class = page["status"] if page["status"] in STATUSES else "active"
        parts.append(
            f'<tr id="screen-{html.escape(page["id"], quote=True)}" data-screen-search="{html.escape(search, quote=True)}">'
            f"<td><strong>{html.escape(page['title'])}</strong><br><span class=path>{html.escape(page['id'])}</span></td>"
            f"<td>{html.escape(table_label(page))}</td>"
            f"<td>{html.escape(page['url'] or '-')}</td>"
            f"<td>{html.escape(page['parent_title'] or '-')}</td>"
            f"<td>{html.escape(page['layout'] or '-')}</td>"
            f"<td>{html.escape(page['exposure'])}</td>"
            f"<td>{html.escape(', '.join(page['roles']))}</td>"
            f"<td>{bool_label(page['desktop'])}</td><td>{bool_label(page['mobile'])}</td>"
            f'<td><span class="status {status_class}">{html.escape(page["status"])}</span></td>'
            f"<td>{html.escape(page['screenshot'] or '-')}</td>"
            "</tr>"
        )
    parts += ["</tbody></table></div></section>", '<section class="section"><h2>API only</h2>']
    if api_only:
        parts += ["<div class=table-wrap><table><thead><tr><th>Method</th><th>API 경로</th><th>용도</th><th>호출 주체</th><th>화면 없음 사유</th></tr></thead><tbody>"]
        for api in api_only:
            parts.append(
                f"<tr><td>{html.escape(api['method'])}</td><td>{html.escape(api['path'])}</td><td>{html.escape(api['purpose'] or '-')}</td><td>{html.escape(api['caller'] or '-')}</td><td>{html.escape(api['reason'] or '-')}</td></tr>"
            )
        parts.append("</tbody></table></div>")
    else:
        parts.append("<p>API only 없음</p>")
    parts += ['</section><section class="section links"><h2>표준 근거</h2>']
    if standards:
        parts += ["<div class=table-wrap><table><thead><tr><th>구분</th><th>문서</th><th>적용 항목</th><th>내용</th><th>링크</th></tr></thead><tbody>"]
        for standard in standards:
            link = f'<a href="{html.escape(standard["url"], quote=True)}" rel="noopener noreferrer">{html.escape(standard["name"])}</a>' if standard["url"] else html.escape(standard["name"])
            parts.append(
                f"<tr><td>{html.escape(standard['source'] or '-')}</td><td>{link}</td><td>{html.escape(standard['scope'] or '-')}</td><td>{html.escape(standard['rule'] or '-')}</td><td>{html.escape(standard['url'] or '-')}</td></tr>"
            )
        parts.append("</tbody></table></div>")
    else:
        parts.append("<p>표준 근거 없음</p>")
    parts += ["</section>", '<section class="section warnings"><h2>검토 메모</h2>']
    if warnings:
        parts.append("<ul>")
        for warning in warnings:
            parts.append(
                f"<li><strong>{html.escape(warning['category'])}</strong>: {html.escape(warning['message'])}"
                + (f" — {html.escape(warning['action'])}" if warning["action"] else "")
                + "</li>"
            )
        parts.append("</ul>")
    else:
        parts.append("<p>검토 메모 없음</p>")
    parts += [
        "</section>",
        '<section class="section"><h2>검토 요청</h2><ol>',
        review_questions,
        "</ol></section>",
        "</main>",
        f"<footer>생성 기준: {html.escape(meta['source'] or '-')} · 적용 기준: {html.escape(meta['standard'] or '-')} · Depth {depth_count} (최소 2, KRDS 최대 3 권장)</footer>",
        f"<script>{javascript}</script>",
        "</body>",
        "</html>",
    ]
    return "\n".join(parts) + "\n"


def main():
    args = parse_args()
    meta, pages, depth_count, warnings, standards, api_only = load_input(args.input)
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    md_path = outdir / f"{args.name}.md"
    html_path = outdir / f"{args.name}.html"
    xlsx_path = outdir / f"{args.name}.xlsx"
    if "md" in args.formats:
        md_path.write_text(build_markdown(meta, pages, depth_count, warnings, standards, api_only), encoding="utf-8")
    if "html" in args.formats:
        html_path.write_text(build_html(meta, pages, depth_count, warnings, standards, api_only), encoding="utf-8")
    if "xlsx" in args.formats:
        write_xlsx(xlsx_path, meta, pages, depth_count, warnings, standards, api_only)
    outputs = []
    for key, path in (("md", md_path), ("html", html_path), ("xlsx", xlsx_path)):
        if key in args.formats:
            outputs.append(repo_path(path))
    if args.sources:
        sources = [item.strip() for item in args.sources.split(",") if item.strip()]
    elif meta["sources"]:
        sources = meta["sources"]
    else:
        sources = [repo_path(args.input)]
        for candidate in ("src/app", "src/pages", "src/routes"):
            if (Path.cwd() / candidate).is_dir():
                sources.append(candidate + "/**")
    unique_sources = []
    for item in sources:
        if item not in unique_sources:
            unique_sources.append(item)
    sources = unique_sources
    manifest_path = outdir / f"{args.name}.artifact.json"
    if not args.no_manifest:
        write_manifest(manifest_path, repo_path(xlsx_path if "xlsx" in args.formats else outputs[0] if outputs else outdir), outputs, sources)
    summary = {
        "md": str(md_path) if "md" in args.formats else "",
        "html": str(html_path) if "html" in args.formats else "",
        "xlsx": str(xlsx_path) if "xlsx" in args.formats else "",
        "pages": len(pages),
        "depth": depth_count,
        "api_mappings": api_count(pages),
        "api_only": len(api_only),
        "standards": len(standards),
        "warnings": len(warnings),
        "manifest": "" if args.no_manifest else str(manifest_path),
    }
    print(json.dumps(summary, ensure_ascii=False))
    if args.strict and warnings:
        sys.exit(1)


if __name__ == "__main__":
    main()
