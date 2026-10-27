#!/usr/bin/env python3
"""오픈소스 고지 페이지를 만든다.

입력: 정규화된 라이선스 JSON (SKILL.md 의 입력 형식 참고)
출력: third-party-notices.html, third-party-notices.md,
      검토 항목이 있으면 third-party-notices-review.md,
      --xlsx 를 주면 third-party-notices.xlsx
"""

import argparse
import difflib
import hashlib
import html
import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

STATUS_LABELS = {
    "none": "고지 불필요",
    "notice": "고지 필요",
    "weak": "소스 공개 검토(약)",
    "strong": "소스 공개 검토(강)",
    "network": "네트워크 공개 검토",
    "unknown": "확인 필요",
}
RANK = {"none": 0, "notice": 1, "weak": 2, "strong": 3, "network": 4, "unknown": 5}

POLICY = {
    "0BSD": "none",
    "MIT-0": "none",
    "Unlicense": "none",
    "CC0-1.0": "none",
    "WTFPL": "none",
    "MIT": "notice",
    "ISC": "notice",
    "BSD-2-Clause": "notice",
    "BSD-3-Clause": "notice",
    "Apache-2.0": "notice",
    "Zlib": "notice",
    "BSL-1.0": "notice",
    "Python-2.0": "notice",
    "MPL-2.0": "weak",
    "EPL-2.0": "weak",
    "CDDL-1.0": "weak",
    "LGPL-2.1-only": "weak",
    "LGPL-2.1-or-later": "weak",
    "LGPL-3.0-only": "weak",
    "LGPL-3.0-or-later": "weak",
    "GPL-2.0-only": "strong",
    "GPL-2.0-or-later": "strong",
    "GPL-3.0-only": "strong",
    "GPL-3.0-or-later": "strong",
    "AGPL-3.0-only": "network",
    "AGPL-3.0-or-later": "network",
}

ALIASES = {
    "mit license": "MIT",
    "expat": "MIT",
    "apache 2.0": "Apache-2.0",
    "apache-2.0": "Apache-2.0",
    "apache license 2.0": "Apache-2.0",
    "apache license, version 2.0": "Apache-2.0",
    "asl 2.0": "Apache-2.0",
    "bsd": "BSD-3-Clause",
    "new bsd": "BSD-3-Clause",
    "bsd 3-clause": "BSD-3-Clause",
    "bsd-3-clause": "BSD-3-Clause",
    "bsd 2-clause": "BSD-2-Clause",
    "bsd-2-clause": "BSD-2-Clause",
    "isc license": "ISC",
    "mozilla public license 2.0": "MPL-2.0",
    "mpl 2.0": "MPL-2.0",
    "gpl-3.0": "GPL-3.0-only",
    "gplv3": "GPL-3.0-only",
    "gpl-2.0": "GPL-2.0-only",
    "gplv2": "GPL-2.0-only",
    "lgpl-3.0": "LGPL-3.0-only",
    "lgpl-2.1": "LGPL-2.1-only",
    "agpl-3.0": "AGPL-3.0-only",
}

EXCEL_TEXT_LIMIT = 30_000
UNKNOWN_LICENSE = "Unknown license"
PAGE_TITLE = "Third-Party Notices"
COPYRIGHT_LINE_RE = re.compile(r"^\s*copyright\b|\(c\)|©", re.I)
TEXT_SIMILARITY = 0.97


def parse_args():
    parser = argparse.ArgumentParser(description="오픈소스 고지 페이지 생성")
    parser.add_argument("--input", required=True, help="정규화된 라이선스 JSON 경로")
    parser.add_argument("--outdir", default="docs/licenses", help="출력 디렉터리")
    parser.add_argument("--name", default="third-party-notices", help="출력 파일 기본 이름")
    parser.add_argument("--xlsx", action="store_true", help="검토용 엑셀도 생성 (openpyxl 필요)")
    parser.add_argument("--exclude-dev", action="store_true", help="scope 가 dev 인 패키지 제외")
    parser.add_argument(
        "--fail-on-review",
        action="store_true",
        help="검토 필요 항목이 있으면 종료 코드 1",
    )
    return parser.parse_args()


def canonical_license(value):
    text = (value or "").strip()
    if not text:
        return ""
    key = text.lower()
    if key in ALIASES:
        return ALIASES[key]
    if text in POLICY:
        return text
    for name in POLICY:
        if name.lower() == key:
            return name
    return text


def classify_single(value):
    return POLICY.get(canonical_license(value), "unknown")


def strip_outer_parens(expression):
    if not (expression.startswith("(") and expression.endswith(")")):
        return expression
    depth = 0
    for index, char in enumerate(expression):
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
            if depth == 0 and index != len(expression) - 1:
                return expression
    return expression[1:-1].strip() if depth == 0 else expression


def split_top_level(expression, operator):
    parts = []
    depth = 0
    start = 0
    index = 0
    upper = expression.upper()
    while index < len(expression):
        char = expression[index]
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        elif depth == 0 and upper.startswith(operator, index):
            before_ok = index == 0 or expression[index - 1].isspace()
            after = index + len(operator)
            after_ok = after >= len(expression) or expression[after].isspace()
            if before_ok and after_ok:
                parts.append(expression[start:index].strip())
                start = after
                index = after
                continue
        index += 1
    if parts:
        parts.append(expression[start:].strip())
    return parts


def classify(value):
    expression = (value or "").strip()
    if not expression:
        return "unknown"
    stripped = strip_outer_parens(expression)
    if stripped != expression:
        return classify(stripped)
    or_parts = split_top_level(expression, "OR")
    if len(or_parts) > 1:
        statuses = [classify(part) for part in or_parts]
        if "unknown" not in statuses:
            return min(statuses, key=lambda item: RANK[item])
    and_parts = split_top_level(expression, "AND")
    if len(and_parts) > 1:
        statuses = [classify(part) for part in and_parts]
        if "unknown" not in statuses:
            return max(statuses, key=lambda item: RANK[item])
    return classify_single(expression)


def normalize_packages(raw_packages):
    merged = {}
    for item in raw_packages:
        name = str(item.get("name", "")).replace("\n", " ").strip()
        if not name:
            continue
        version = str(item.get("version", "")).replace("\n", " ").strip()
        key = (name, version)
        license_value = str(item.get("license", "")).replace("\n", " ").strip()
        entry = merged.get(key)
        if entry is None:
            entry = {
                "name": name,
                "version": version,
                "license": license_value,
                "copyright": str(item.get("copyright", "")).strip(),
                "repository": str(item.get("repository", "")).strip(),
                "scope": str(item.get("scope", "unknown")).strip().lower() or "unknown",
                "dependency": str(item.get("dependency", "")).strip(),
                "license_text": str(item.get("license_text", "")),
                "license_source": str(item.get("license_source", "")).strip(),
                "notice_text": str(item.get("notice_text", "")).strip(),
            }
            merged[key] = entry
        else:
            if not entry["license"] and license_value:
                entry["license"] = license_value
            elif license_value and license_value not in entry["license"]:
                entry["license"] = (
                    f"{entry['license']} OR {license_value}" if entry["license"] else license_value
                )
            for field in ("copyright", "repository", "license_source"):
                if not entry[field] and item.get(field):
                    entry[field] = str(item[field]).strip()
            if not entry["license_text"] and item.get("license_text"):
                entry["license_text"] = str(item["license_text"])
            if not entry["notice_text"] and item.get("notice_text"):
                entry["notice_text"] = str(item["notice_text"]).strip()
            if entry["scope"] == "unknown" and item.get("scope") in ("runtime", "dev"):
                entry["scope"] = str(item["scope"])
            if item.get("dependency") == "direct":
                entry["dependency"] = "direct"

    packages = []
    for entry in merged.values():
        entry["license"] = canonical_license(entry["license"])
        entry["status"] = classify(entry["license"])
        packages.append(entry)
    packages.sort(key=lambda item: (item["name"].lower(), item["version"], item["license"].lower()))
    return packages


def review_items(packages):
    items = []
    for package in packages:
        reasons = []
        actions = []
        if package["status"] in ("weak", "strong", "network"):
            reasons.append("copyleft 라이선스, 소스 공개 범위 검토 필요")
            actions.append("법무 검토 후 배포 여부와 소스 공개 범위 결정")
        if package["status"] == "unknown":
            reasons.append("라이선스 확인 필요")
            actions.append("저장소의 LICENSE 원문을 확인해 license 값 보정")
        if re.search(r"\s+(OR|AND)\s+", package["license"], re.I):
            reasons.append("복수 라이선스 표현, 적용 조건 확인 필요")
            actions.append("배포에 적용할 라이선스를 정하고 해당 전문이 수록됐는지 확인")
        if not package["license_text"]:
            reasons.append("라이선스 전문 없음")
            actions.append("저장소에서 LICENSE 원문을 받아 license_text 채우기")
        if package["scope"] == "dev":
            reasons.append("개발 의존성, 배포 대상 여부 확인 필요")
            actions.append("앱 산출물에 포함되지 않으면 packages 에서 제외")
        if reasons:
            items.append(
                {
                    "package": package,
                    "reason": " / ".join(reasons),
                    "action": " / ".join(actions),
                }
            )
    items.sort(
        key=lambda item: (
            -RANK[item["package"]["status"]],
            item["package"]["name"].lower(),
        )
    )
    return items


def group_notices(packages):
    """라이선스별로 묶고, 표기 차이만 있는 같은 전문은 한 번만 수록한다."""
    groups = {}
    for package in packages:
        name = package["license"] or UNKNOWN_LICENSE
        group = groups.setdefault(name, {"license": name, "variants": {}, "missing": []})
        if not package["license_text"]:
            group["missing"].append(package)
            continue
        normalized = normalized_words(package["license_text"])
        match = None
        best = 0.0
        for key, variant in group["variants"].items():
            matcher = difflib.SequenceMatcher(None, normalized, variant["normalized"])
            if matcher.quick_ratio() < TEXT_SIMILARITY:
                continue
            ratio = matcher.ratio()
            if ratio > best:
                best = ratio
                match = key
        if match is not None and best >= TEXT_SIMILARITY:
            group["variants"][match]["packages"].append(package)
            continue
        digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:12]
        group["variants"][digest] = {
            "digest": digest,
            "text": package["license_text"],
            "normalized": normalized,
            "packages": [package],
        }

    ordered = []
    for name in sorted(groups, key=lambda value: (value == UNKNOWN_LICENSE, value.lower())):
        group = groups[name]
        variants = sorted(
            group["variants"].values(),
            key=lambda variant: variant["packages"][0]["name"].lower(),
        )
        ordered.append(
            {
                "license": group["license"],
                "variants": variants,
                "missing": sorted(group["missing"], key=lambda p: p["name"].lower()),
            }
        )
    return ordered


def normalized_words(text):
    """저작권 줄과 줄바꿈·따옴표 차이를 지워 전문의 뼈대를 비교한다."""
    lines = [line for line in text.splitlines() if not COPYRIGHT_LINE_RE.search(line)]
    lowered = "\n".join(lines).lower().replace("licence", "license")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", lowered)).strip()


def slugify(value):
    slug = re.sub(r"[^a-z0-9]+", "-", (value or "").lower()).strip("-")
    return slug or "license"


def license_slugs(groups):
    seen = {}
    slugs = {}
    for group in groups:
        base = slugify(group["license"])
        seen[base] = seen.get(base, 0) + 1
        slugs[group["license"]] = base if seen[base] == 1 else f"{base}-{seen[base]}"
    return slugs


def package_label(package):
    return f"{package['name']} {package['version']}".strip()


def copyright_line(package):
    return " · ".join(
        line.strip() for line in package.get("copyright", "").splitlines() if line.strip()
    )


def fence(text):
    longest = max((len(match.group(0)) for match in re.finditer(r"`+", text)), default=0)
    marker = "`" * max(3, longest + 1)
    return f"{marker}\n{text.rstrip()}\n{marker}"


def build_markdown(data, groups):
    project = str(data.get("project", ""))
    version = str(data.get("app_version", ""))
    generated_at = str(data.get("generated_at", ""))
    lines = [f"# {PAGE_TITLE}", ""]
    title = " ".join(part for part in (project, version) if part)
    if title:
        lines.extend([f"**{title}**", ""])
    if generated_at:
        lines.extend([f"Generated: {str(generated_at)[:10]}", ""])
    lines.extend(
        [
            "This product includes the third-party software listed below. "
            "Each license text and copyright notice is reproduced for attribution. "
            "Where libraries share the same license, the text appears once and each "
            "copyright notice is listed with its library.",
            "",
        ]
    )
    for group in groups:
        lines.extend([f"## {group['license']}", ""])
        for variant in group["variants"]:
            lines.extend(["Applies to:", ""])
            for package in variant["packages"]:
                line = f"- {package_label(package)}"
                notice = copyright_line(package)
                if notice:
                    line += f" — {notice}"
                lines.append(line)
            lines.extend(["", "License text:", "", fence(variant["text"]), ""])
            seen_notices = []
            for package in variant["packages"]:
                notice_text = package.get("notice_text", "").strip()
                if notice_text and notice_text not in seen_notices:
                    seen_notices.append(notice_text)
            for notice_text in seen_notices:
                lines.extend(["NOTICE:", "", fence(notice_text), ""])
        if group["missing"]:
            lines.extend(["License text not found for:", ""])
            for package in group["missing"]:
                lines.append(f"- {package_label(package)}")
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def build_html(data, groups):
    project = str(data.get("project", ""))
    version = str(data.get("app_version", ""))
    generated_at = str(data.get("generated_at", ""))[:10]
    app_title = " ".join(part for part in (project, version) if part)
    document_title = f"{PAGE_TITLE} — {app_title}" if app_title else PAGE_TITLE
    package_count = sum(
        len(variant["packages"]) for group in groups for variant in group["variants"]
    ) + sum(len(group["missing"]) for group in groups)
    license_count = len(groups)
    slugs = license_slugs(groups)
    initial_count = f"{package_count} libraries · {license_count} licenses"
    meta_parts = [app_title] if app_title else []
    if generated_at:
        meta_parts.append(f"Generated {generated_at}")
    meta_parts.append(f"{package_count} libraries")
    meta_line = " · ".join(meta_parts)

    css = """
:root {
  color-scheme: light dark;
  --bg: #ffffff;
  --fg: #17181c;
  --muted: #5c6270;
  --line: #e2e5ea;
  --panel: #f6f7f9;
  --accent: #1f3864;
  --code: #f3f4f6;
}
@media (prefers-color-scheme: dark) {
  :root {
    --bg: #111216;
    --fg: #e8eaee;
    --muted: #9aa1ae;
    --line: #2a2d35;
    --panel: #181a20;
    --accent: #9db8e8;
    --code: #1a1c22;
  }
}
* { box-sizing: border-box; }
html { -webkit-text-size-adjust: 100%; }
body {
  margin: 0;
  background: var(--bg);
  color: var(--fg);
  font: 16px/1.6 -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto,
    "Hiragino Sans", "Noto Sans KR", sans-serif;
}
main, header, footer { max-width: 60rem; margin: 0 auto; padding: 0 1.25rem; }
header { padding-top: 2.5rem; padding-bottom: 1rem; }
h1 { font-size: 1.75rem; line-height: 1.25; margin: 0 0 .5rem; letter-spacing: -.01em; }
h2 { font-size: 1.125rem; margin: 0; }
p { margin: .5rem 0; }
.eyebrow { color: var(--accent); font-size: .75rem; font-weight: 700; letter-spacing: .08em; text-transform: uppercase; }
.meta { color: var(--muted); font-size: .875rem; }
.lede { margin-top: 1rem; max-width: 44rem; }
.controls {
  display: flex; flex-wrap: wrap; align-items: center; gap: .75rem;
  margin: 1.5rem 0 1rem; padding: 1rem; background: var(--panel);
  border: 1px solid var(--line); border-radius: .75rem;
}
.controls label { font-weight: 600; font-size: .875rem; }
.controls input {
  flex: 1 1 16rem; min-width: 0; padding: .55rem .75rem;
  color: var(--fg); background: var(--bg);
  border: 1px solid var(--line); border-radius: .5rem; font: inherit;
}
.controls input:focus-visible { outline: 2px solid var(--accent); outline-offset: 1px; }
.count { margin: 0; color: var(--muted); font-size: .8125rem; }
.toc { margin: 1.5rem 0 2rem; }
.toc h2 { font-size: .8125rem; color: var(--muted); letter-spacing: .06em; text-transform: uppercase; margin-bottom: .5rem; }
.toc ul { display: flex; flex-wrap: wrap; gap: .5rem; margin: 0; padding: 0; list-style: none; }
.toc a {
  display: inline-flex; gap: .4rem; align-items: baseline;
  padding: .3rem .6rem; color: var(--fg); text-decoration: none;
  background: var(--panel); border: 1px solid var(--line); border-radius: 999px;
  font-size: .8125rem;
}
.toc a:hover { border-color: var(--accent); }
.toc .toc-count { color: var(--muted); font-size: .75rem; }
.license { padding: 1.5rem 0; border-top: 1px solid var(--line); scroll-margin-top: 1rem; }
.license-head { display: flex; align-items: baseline; gap: .6rem; margin-bottom: .75rem; }
.badge {
  padding: .05rem .5rem; color: var(--muted); font-size: .75rem;
  border: 1px solid var(--line); border-radius: 999px;
}
.packages { margin: 0 0 1rem; padding-left: 1.1rem; }
.applies { margin: 0 0 .3rem; color: var(--muted); font-size: .75rem; letter-spacing: .06em; text-transform: uppercase; }
.packages li { margin: .3rem 0; }
.packages .name { font-weight: 600; }
.packages .version { color: var(--muted); font-size: .8125rem; margin-left: .35rem; }
.packages .copyright { display: block; color: var(--muted); font-size: .8125rem; }
details { margin: .75rem 0 0; }
summary {
  cursor: pointer; font-weight: 600; font-size: .875rem;
  padding: .45rem .65rem; background: var(--panel);
  border: 1px solid var(--line); border-radius: .5rem;
}
summary:hover { border-color: var(--accent); }
pre {
  margin: .5rem 0 0; padding: 1rem; overflow: auto;
  white-space: pre-wrap; word-break: break-word;
  background: var(--code); border: 1px solid var(--line); border-radius: .5rem;
  font: .8125rem/1.5 ui-monospace, SFMono-Regular, Menlo, Consolas, monospace;
}
.notice-label { margin: 1rem 0 0; font-weight: 600; font-size: .875rem; }
.variant + .variant { margin-top: 1.25rem; }
.missing { margin-top: 1rem; padding: .75rem 1rem; border: 1px dashed var(--line); border-radius: .5rem; }
.missing p { color: var(--muted); font-size: .875rem; }
.missing ul { margin: .25rem 0 0; padding-left: 1.1rem; font-size: .875rem; }
footer { padding-top: 2rem; padding-bottom: 3rem; color: var(--muted); font-size: .8125rem; border-top: 1px solid var(--line); }
[hidden] { display: none !important; }
@media print {
  .controls, .toc, footer { display: none; }
  details { display: block; }
  details > summary { display: none; }
  details > pre { display: block !important; }
  body { font-size: 11pt; }
  .license { break-inside: avoid-page; }
}
""".strip()

    javascript = """
(function () {
  var input = document.getElementById("search");
  if (!input) return;
  var count = document.getElementById("result-count");
  var sections = Array.prototype.slice.call(document.querySelectorAll("section.license"));
  var links = Array.prototype.slice.call(document.querySelectorAll("nav.toc a"));
  var allItems = Array.prototype.slice.call(document.querySelectorAll("li.package"));
  var initial = count ? count.textContent : "";
  function normalize(value) { return (value || "").toLowerCase(); }
  function apply() {
    var tokens = normalize(input.value).trim().split(/\\s+/).filter(Boolean);
    var visible = 0;
    sections.forEach(function (section) {
      var groupMatch = tokens.length > 0 && tokens.every(function (token) {
        return normalize(section.getAttribute("data-license")).indexOf(token) !== -1;
      });
      var sectionShown = 0;
      var variants = Array.prototype.slice.call(section.querySelectorAll(".variant"));
      variants.forEach(function (variant) {
        var items = Array.prototype.slice.call(variant.querySelectorAll("li.package"));
        var shown = 0;
        items.forEach(function (item) {
          var match = groupMatch || tokens.every(function (token) {
            return normalize(item.getAttribute("data-search")).indexOf(token) !== -1;
          });
          item.hidden = !match;
          if (match) shown += 1;
        });
        variant.hidden = shown === 0;
        var details = variant.querySelectorAll("details");
        for (var i = 0; i < details.length; i += 1) {
          details[i].open = tokens.length > 0 && shown > 0;
        }
        sectionShown += shown;
      });
      section.hidden = sectionShown === 0;
      var badge = section.querySelector(".badge");
      if (badge) {
        badge.textContent = tokens.length === 0 ? badge.getAttribute("data-total") : sectionShown;
      }
      visible += sectionShown;
    });
    links.forEach(function (link) {
      var target = document.querySelector(link.getAttribute("href"));
      var row = link.parentNode;
      row.hidden = !!(target && target.hidden);
      var countLabel = link.querySelector(".toc-count");
      if (countLabel && target) {
        countLabel.textContent = tokens.length === 0
          ? link.getAttribute("data-total")
          : target.querySelectorAll("li.package:not([hidden])").length;
      }
    });
    if (count) {
      count.textContent = tokens.length === 0
        ? initial
        : visible + " of " + allItems.length + " libraries shown";
    }
  }
  input.addEventListener("input", apply);
})();
""".strip()

    parts = [
        "<!doctype html>",
        '<html lang="en">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        '<meta http-equiv="Content-Security-Policy" '
        "content=\"default-src 'none'; style-src 'unsafe-inline'; script-src 'unsafe-inline'\">",
        f"<title>{html.escape(document_title)}</title>",
        f"<style>\n{css}\n</style>",
        "</head>",
        "<body>",
        "<header>",
        '<p class="eyebrow">Open source notices</p>',
        f"<h1>{PAGE_TITLE}</h1>",
        f'<p class="meta">{html.escape(meta_line)}</p>',
        '<p class="lede">This product includes the third-party software listed below. '
        "Each license text and copyright notice is reproduced for attribution. "
        "Where libraries share the same license, the text appears once and each "
        "copyright notice is listed with its library.</p>",
        "</header>",
        "<main>",
        '<div class="controls">',
        '<label for="search">Search libraries</label>',
        '<input id="search" type="search" placeholder="Library, version, or license" '
        'autocomplete="off" spellcheck="false">',
        f'<p class="count" id="result-count" aria-live="polite">{html.escape(initial_count)}</p>',
        "</div>",
        "<noscript><p class=\"count\">Search needs JavaScript. All licenses are listed below.</p></noscript>",
        '<nav class="toc" aria-label="License index">',
        "<h2>Licenses</h2>",
        "<ul>",
    ]
    for group in groups:
        slug = slugs[group["license"]]
        count = sum(len(variant["packages"]) for variant in group["variants"]) + len(
            group["missing"]
        )
        parts.append(
            f'<li><a href="#license-{slug}" data-total="{count}">{html.escape(group["license"])}'
            f'<span class="toc-count">{count}</span></a></li>'
        )
    parts.extend(["</ul>", "</nav>"])

    for group in groups:
        slug = slugs[group["license"]]
        count = sum(len(variant["packages"]) for variant in group["variants"]) + len(
            group["missing"]
        )
        parts.append(
            f'<section class="license" id="license-{slug}" '
            f'data-license="{html.escape(group["license"], quote=True)}">'
        )
        parts.append(
            f'<div class="license-head"><h2>{html.escape(group["license"])}</h2>'
            f'<span class="badge" data-total="{count}">{count}</span></div>'
        )
        for index, variant in enumerate(group["variants"], start=1):
            parts.append('<div class="variant">')
            parts.append('<p class="applies">Applies to</p>')
            parts.append('<ul class="packages">')
            for package in variant["packages"]:
                search_value = " ".join(
                    part
                    for part in (
                        package["name"],
                        package["version"],
                        group["license"],
                        copyright_line(package),
                    )
                    if part
                )
                parts.append(
                    f'<li class="package" data-search="{html.escape(search_value.lower(), quote=True)}">'
                    f'<span class="name">{html.escape(package["name"])}</span>'
                )
                if package["version"]:
                    parts.append(
                        f'<span class="version">{html.escape(package["version"])}</span>'
                    )
                notice = copyright_line(package)
                if notice:
                    parts.append(f'<span class="copyright">{html.escape(notice)}</span>')
                parts.append("</li>")
            parts.append("</ul>")
            label = "License text" if len(group["variants"]) == 1 else f"License text {index}"
            parts.append(
                f"<details><summary>{label}</summary>"
                f"<pre tabindex=\"0\">{html.escape(variant['text'].rstrip())}</pre>"
            )
            seen_notices = []
            for package in variant["packages"]:
                notice_text = package.get("notice_text", "").strip()
                if notice_text and notice_text not in seen_notices:
                    seen_notices.append(notice_text)
            for notice_text in seen_notices:
                parts.append(
                    '<p class="notice-label">NOTICE</p>'
                    f"<pre tabindex=\"0\">{html.escape(notice_text.rstrip())}</pre>"
                )
            parts.append("</details>")
            parts.append("</div>")
        if group["missing"]:
            parts.append('<div class="variant missing"><p>License text not found for:</p><ul>')
            for package in group["missing"]:
                parts.append(
                    f'<li class="package" data-search="{html.escape((package_label(package) + " " + group["license"]).lower(), quote=True)}">'
                    f"{html.escape(package_label(package))}</li>"
                )
            parts.append("</ul></div>")
        parts.append("</section>")

    parts.extend(
        [
            "</main>",
            "<footer>Generated from the project SBOM. License texts are reproduced verbatim "
            "from the package sources.</footer>",
            f"<script>\n{javascript}\n</script>",
            "</body>",
            "</html>",
        ]
    )
    return "\n".join(parts) + "\n"


def build_review_markdown(data, review):
    lines = [
        "# 고지 검토 필요",
        "",
        f"- 프로젝트: {data.get('project', '')}",
        f"- 앱 버전: {data.get('app_version', '')}",
        f"- 생성 시각: {data.get('generated_at', '')}",
        f"- 검토 항목: {len(review)}",
        "",
        "이 파일은 내부 검토용입니다. 앱이나 배포 페이지에 포함하지 않습니다.",
        "",
        "| 패키지 | 버전 | 라이선스 | 사유 | 조치 |",
        "| --- | --- | --- | --- | --- |",
    ]
    for item in review:
        package = item["package"]
        lines.append(
            "| "
            + " | ".join(
                [
                    str(package["name"]).replace("|", "\\|"),
                    str(package["version"]).replace("|", "\\|"),
                    (package["license"] or "unknown").replace("|", "\\|"),
                    item["reason"].replace("|", "\\|"),
                    item["action"].replace("|", "\\|"),
                ]
            )
            + " |"
        )
    return "\n".join(lines) + "\n"


def write_workbook(path, data, packages, review, groups, name):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit(
            "엑셀을 만들려면 openpyxl 이 필요합니다. "
            "'python3 -m pip install --user openpyxl' 실행 후 다시 시도하세요. "
            "HTML 과 md 는 openpyxl 없이도 생성됩니다."
        )

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(size=14, bold=True)
    section_font = Font(bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    status_fills = {
        "none": PatternFill("solid", fgColor="D9D9D9"),
        "notice": PatternFill("solid", fgColor="C6E0B4"),
        "weak": PatternFill("solid", fgColor="FFE699"),
        "strong": PatternFill("solid", fgColor="F4B183"),
        "network": PatternFill("solid", fgColor="F8CBAD"),
        "unknown": PatternFill("solid", fgColor="E4DFEC"),
    }

    def style_header(ws, widths):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A2"

    def write_summary(ws):
        ws["A1"] = "오픈소스 라이선스 고지"
        ws["A1"].font = title_font
        ws.merge_cells("A1:D1")
        scope_counts = Counter(p["scope"] for p in packages)
        status_counts = Counter(p["status"] for p in packages)
        rows = [
            ("프로젝트", data.get("project", "")),
            ("앱 버전", data.get("app_version", "")),
            ("생성 시각", data.get("generated_at", "")),
            ("대상 패키지", len(packages)),
            ("런타임", scope_counts.get("runtime", 0)),
            ("개발", scope_counts.get("dev", 0)),
            ("확인 필요(범위)", scope_counts.get("unknown", 0)),
            ("검토 필요 항목", len(review)),
            ("라이선스 전문 없음", sum(1 for p in packages if not p["license_text"])),
        ]
        for index, (label, value) in enumerate(rows, start=3):
            ws.cell(row=index, column=1, value=label).font = section_font
            ws.cell(row=index, column=2, value=value)
        start = len(rows) + 5
        ws.cell(row=start, column=1, value="고지 의무 집계").font = section_font
        for offset, status in enumerate(STATUS_LABELS, start=1):
            row = start + offset
            cell = ws.cell(row=row, column=1, value=STATUS_LABELS[status])
            cell.fill = status_fills[status]
            cell.border = border
            ws.cell(row=row, column=2, value=status_counts.get(status, 0)).border = border
        note_row = start + len(STATUS_LABELS) + 3
        ws.cell(row=note_row, column=1, value="생성 규칙").font = section_font
        ws.cell(
            row=note_row + 1,
            column=1,
            value=f"라이선스 전문은 {name}.html 과 {name}.md 에 원문 그대로 수록",
        )
        ws.cell(row=note_row + 2, column=1, value="copyleft 와 라이선스 미확인 항목은 검토 필요 시트에서 확인")
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 48

    def write_packages(ws):
        headers = [
            "패키지",
            "버전",
            "라이선스",
            "고지 의무",
            "저작권 표시",
            "범위",
            "의존성",
            "전문",
            "라이선스 파일",
            "저장소",
        ]
        ws.append(headers)
        style_header(ws, [26, 12, 22, 18, 46, 10, 10, 8, 48, 44])
        for package in packages:
            ws.append(
                [
                    package["name"],
                    package["version"],
                    package["license"] or "unknown",
                    STATUS_LABELS[package["status"]],
                    package["copyright"],
                    package["scope"],
                    package["dependency"],
                    "있음" if package["license_text"] else "없음",
                    package["license_source"],
                    package["repository"],
                ]
            )
            row = ws.max_row
            for cell in ws[row]:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=row, column=4).fill = status_fills[package["status"]]
            if not package["license_text"]:
                ws.cell(row=row, column=8).fill = PatternFill("solid", fgColor="F8CBAD")

    def write_license_texts(ws):
        headers = ["라이선스", "패키지 수", "글자 수", "전문", "비고"]
        ws.append(headers)
        style_header(ws, [24, 10, 10, 110, 40])
        for group in groups:
            for variant in group["variants"]:
                text = variant["text"]
                truncated = len(text) > EXCEL_TEXT_LIMIT
                shown = text[:EXCEL_TEXT_LIMIT]
                note = "전문 생략, html 파일 참조" if truncated else ""
                ws.append(
                    [group["license"], len(variant["packages"]), len(text), shown, note]
                )
                row = ws.max_row
                for cell in ws[row]:
                    cell.border = border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    def write_review(ws):
        headers = ["패키지", "버전", "라이선스", "사유", "조치"]
        ws.append(headers)
        style_header(ws, [26, 12, 22, 52, 52])
        for item in review:
            package = item["package"]
            ws.append(
                [
                    package["name"],
                    package["version"],
                    package["license"] or "unknown",
                    item["reason"],
                    item["action"],
                ]
            )
            row = ws.max_row
            for cell in ws[row]:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    def write_logs(ws):
        headers = ["도구", "명령", "상태", "비고"]
        ws.append(headers)
        style_header(ws, [22, 90, 12, 40])
        for log in data.get("logs", []):
            ws.append(
                [
                    log.get("tool", ""),
                    log.get("command", ""),
                    log.get("status", ""),
                    log.get("note", ""),
                ]
            )
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=4):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "요약"
    write_summary(summary)
    write_packages(workbook.create_sheet("패키지"))
    write_license_texts(workbook.create_sheet("라이선스 전문"))
    write_review(workbook.create_sheet("검토 필요"))
    write_logs(workbook.create_sheet("실행 로그"))
    workbook.save(path)


def main():
    args = parse_args()
    input_path = Path(args.input)
    if not input_path.is_file():
        sys.exit(f"입력 파일이 없습니다: {input_path}")
    try:
        data = json.loads(input_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        sys.exit(f"입력 JSON 을 읽지 못했습니다: {error}")
    raw_packages = data.get("packages")
    if not isinstance(raw_packages, list) or not raw_packages:
        sys.exit("입력 JSON 의 packages 배열이 비어 있습니다.")
    if args.exclude_dev:
        raw_packages = [
            item for item in raw_packages if str(item.get("scope", "")).strip().lower() != "dev"
        ]
        if not raw_packages:
            sys.exit("--exclude-dev 적용 후 대상 패키지가 없습니다.")

    packages = normalize_packages(raw_packages)
    if not packages:
        sys.exit("정규화된 패키지가 없습니다. 입력 형식을 확인하세요.")
    data.setdefault("generated_at", datetime.now().astimezone().isoformat(timespec="seconds"))
    review = review_items(packages)
    groups = group_notices(packages)

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    html_path = outdir / f"{args.name}.html"
    md_path = outdir / f"{args.name}.md"
    review_path = outdir / f"{args.name}-review.md"
    xlsx_path = outdir / f"{args.name}.xlsx"

    html_path.write_text(build_html(data, groups), encoding="utf-8")
    md_path.write_text(build_markdown(data, groups), encoding="utf-8")
    if review:
        review_path.write_text(build_review_markdown(data, review), encoding="utf-8")
    elif review_path.exists():
        review_path.unlink()
    if args.xlsx:
        write_workbook(xlsx_path, data, packages, review, groups, args.name)

    summary = {
        "html": str(html_path),
        "md": str(md_path),
        "review": str(review_path) if review else "",
        "xlsx": str(xlsx_path) if args.xlsx else "",
        "packages": len(packages),
        "licenses": len(groups),
        "license_texts": sum(len(group["variants"]) for group in groups),
        "review_items": len(review),
        "status": {
            STATUS_LABELS[key]: value
            for key, value in Counter(p["status"] for p in packages).items()
        },
    }
    print(json.dumps(summary, ensure_ascii=False))

    if args.fail_on_review and review:
        sys.exit(1)


if __name__ == "__main__":
    main()
