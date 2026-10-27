#!/usr/bin/env python3
"""브랜치 diff 점검 항목을 릴리스 점검 엑셀 보고서로 변환한다.

입력: 정규화된 release-check JSON (SKILL.md 의 입력 JSON 형식 참고)
출력: release-check-YYYY-MM-DD.xlsx
"""

import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit(
        "openpyxl 이 필요합니다. 'python3 -m pip install --user openpyxl' 실행 후 다시 시도하세요."
    )

PRIORITIES = ["P0", "P1", "P2"]
AUDIENCES = ["qa", "dev"]
AUDIENCE_LABELS = {"qa": "QA", "dev": "개발자"}
RESULTS = ["미확인", "통과", "실패", "해당 없음"]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True)
SECTION_FONT = Font(bold=True)
SEPARATOR_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRIORITY_FILLS = {
    "P0": PatternFill("solid", fgColor="F8CBAD"),
    "P1": PatternFill("solid", fgColor="FFE699"),
    "P2": PatternFill("solid", fgColor="E2EFDA"),
}
RESULT_FILLS = {
    "미확인": PatternFill("solid", fgColor="FFF2CC"),
    "통과": PatternFill("solid", fgColor="C6E0B4"),
    "실패": PatternFill("solid", fgColor="F4B183"),
    "해당 없음": PatternFill("solid", fgColor="D9D9D9"),
}


def parse_args():
    parser = argparse.ArgumentParser(description="릴리스 점검 엑셀 보고서 생성")
    parser.add_argument("--input", required=True, help="정규화된 release-check JSON 경로")
    parser.add_argument("--outdir", default="docs/qa", help="출력 디렉터리")
    parser.add_argument("--date", help="파일명에 쓸 날짜 (기본: 오늘, YYYY-MM-DD)")
    parser.add_argument("--name", help="출력 파일 기본 이름 (기본: release-check-YYYY-MM-DD)")
    return parser.parse_args()


def load_input(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    items = data.get("items")
    if not isinstance(items, list) or not items:
        sys.exit("입력 JSON 의 items 배열이 비어 있습니다.")
    return data


def normalize_items(items):
    normalized = []
    for item in items:
        priority = str(item.get("priority", "")).upper()
        audience = str(item.get("audience", "")).lower()
        if priority not in PRIORITIES:
            sys.exit(f"priority 는 P0/P1/P2 중 하나여야 합니다: {item!r}")
        if audience not in AUDIENCES:
            sys.exit(f"audience 는 qa/dev 중 하나여야 합니다: {item!r}")
        result = str(item.get("status") or "미확인").strip()
        if result not in RESULTS:
            result = "미확인"
        normalized.append(
            {
                "priority": priority,
                "audience": audience,
                "area": str(item.get("area", "")).strip(),
                "item": str(item.get("item", "")).strip(),
                "expected": str(item.get("expected", "")).strip(),
                "evidence": str(item.get("evidence", "")).strip(),
                "status": result,
                "owner": str(item.get("owner", "")).strip(),
                "checked_at": str(item.get("checked_at", "")).strip(),
                "note": str(item.get("note", "")).strip(),
            }
        )
    normalized.sort(
        key=lambda i: (
            0 if i["audience"] == "qa" else 1,
            PRIORITIES.index(i["priority"]),
            i["area"],
            i["item"],
        )
    )
    counters = Counter()
    for item in normalized:
        key = (item["priority"], item["audience"])
        counters[key] += 1
        label = "DEV" if item["audience"] == "dev" else "QA"
        item["id"] = f"{item['priority']}-{label}-{counters[key]:02d}"
    return normalized


def style_header(ws, columns):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    for index, width in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"


def write_summary(ws, data, items, generated_at):
    ws["A1"] = "릴리스 전 브랜치 diff 점검 보고서"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    base = data.get("base", "")
    head = data.get("head", "")
    stats = data.get("stats", {})
    rows = [
        ("프로젝트", data.get("project", "")),
        ("배포 버전", data.get("release_version", "")),
        ("비교 브랜치", f"{base} → {head}"),
        ("merge-base", data.get("merge_base", "")),
        ("생성 시각", data.get("generated_at", generated_at)),
        ("커밋 수", stats.get("commits", "")),
        ("변경 파일 수", stats.get("files", "")),
        ("증감 라인", f"+{stats.get('additions', 0)} / -{stats.get('deletions', 0)}"),
        ("점검 항목", len(items)),
    ]
    for index, (label, value) in enumerate(rows, start=3):
        ws.cell(row=index, column=1, value=label).font = SECTION_FONT
        ws.cell(row=index, column=2, value=value)

    priority_counts = Counter(i["priority"] for i in items)
    audience_counts = Counter(i["audience"] for i in items)
    start = 14
    ws.cell(row=start, column=1, value="우선순위 집계").font = SECTION_FONT
    headers = ["우선순위", "QA", "개발자", "합계"]
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=start + 1, column=col, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    for offset, priority in enumerate(PRIORITIES, start=2):
        qa = sum(1 for i in items if i["priority"] == priority and i["audience"] == "qa")
        dev = sum(1 for i in items if i["priority"] == priority and i["audience"] == "dev")
        row = start + offset
        ws.cell(row=row, column=1, value=priority).border = BORDER
        ws.cell(row=row, column=1).fill = PRIORITY_FILLS[priority]
        ws.cell(row=row, column=2, value=qa).border = BORDER
        ws.cell(row=row, column=3, value=dev).border = BORDER
        ws.cell(row=row, column=4, value=qa + dev).border = BORDER
    total_row = start + 2 + len(PRIORITIES)
    ws.cell(row=total_row, column=1, value="합계").font = SECTION_FONT
    ws.cell(row=total_row, column=2, value=audience_counts.get("qa", 0))
    ws.cell(row=total_row, column=3, value=audience_counts.get("dev", 0))
    ws.cell(row=total_row, column=4, value=len(items))

    start = total_row + 3
    ws.cell(row=start, column=1, value="확인 결과 집계").font = SECTION_FONT
    for offset, result in enumerate(RESULTS, start=1):
        ws.cell(row=start + offset, column=1, value=result).border = BORDER
        ws.cell(row=start + offset, column=1).fill = RESULT_FILLS[result]
        ws.cell(
            row=start + offset,
            column=2,
            value=sum(1 for i in items if i["status"] == result),
        ).border = BORDER

    note_row = start + len(RESULTS) + 3
    ws.cell(row=note_row, column=1, value="배치 규칙").font = SECTION_FONT
    ws.cell(row=note_row + 1, column=1, value="QA 항목 아래에 개발자 전용 항목을 배치합니다.")
    ws.cell(row=note_row + 2, column=1, value="확인 결과는 미확인/통과/실패/해당 없음 중 하나로 기록합니다.")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14


def write_items(ws, items):
    headers = [
        "ID",
        "우선순위",
        "영역",
        "점검 항목",
        "기대 결과",
        "변경 근거",
        "구분",
        "확인 결과",
        "확인자",
        "확인일",
        "비고",
    ]
    ws.append(headers)
    style_header(ws, [18, 10, 18, 52, 40, 42, 10, 12, 12, 12, 30])

    def append_item(item):
        ws.append(
            [
                item["id"],
                item["priority"],
                item["area"],
                item["item"],
                item["expected"],
                item["evidence"],
                AUDIENCE_LABELS[item["audience"]],
                item["status"],
                item["owner"],
                item["checked_at"],
                item["note"],
            ]
        )
        row = ws.max_row
        for cell in ws[row]:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row, column=2).fill = PRIORITY_FILLS[item["priority"]]
        ws.cell(row=row, column=8).fill = RESULT_FILLS[item["status"]]
        ws.cell(row=row, column=1).font = Font(bold=True)

    qa_items = [i for i in items if i["audience"] == "qa"]
    dev_items = [i for i in items if i["audience"] == "dev"]
    for item in qa_items:
        append_item(item)
    if dev_items:
        ws.append(["개발자 전용 점검 (코드·설정·DB 확인 필요)"] + [""] * 10)
        separator = ws.max_row
        ws.merge_cells(start_row=separator, start_column=1, end_row=separator, end_column=11)
        cell = ws.cell(row=separator, column=1)
        cell.fill = SEPARATOR_FILL
        cell.font = SECTION_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        for column in range(1, 12):
            ws.cell(row=separator, column=column).border = BORDER
        for item in dev_items:
            append_item(item)

    priority_validation = DataValidation(
        type="list", formula1='"' + ",".join(PRIORITIES) + '"', allow_blank=False
    )
    result_validation = DataValidation(
        type="list", formula1='"' + ",".join(RESULTS) + '"', allow_blank=False
    )
    ws.add_data_validation(priority_validation)
    ws.add_data_validation(result_validation)
    qa_end = 1 + len(qa_items)
    dev_start = qa_end + 2 if dev_items else qa_end + 1
    last_row = ws.max_row
    ranges = []
    if qa_items:
        ranges += [f"B2:B{qa_end}", f"H2:H{qa_end}"]
    if dev_items:
        ranges += [f"B{dev_start}:B{last_row}", f"H{dev_start}:H{last_row}"]
    for range_ref in ranges:
        target = priority_validation if range_ref.startswith("B") else result_validation
        target.add(range_ref)


def write_files(ws, files):
    headers = ["경로", "상태", "증감", "영역", "관련 항목 ID"]
    ws.append(headers)
    style_header(ws, [70, 8, 14, 18, 24])
    for entry in files:
        additions = entry.get("additions", 0)
        deletions = entry.get("deletions", 0)
        ws.append(
            [
                entry.get("path", ""),
                entry.get("status", ""),
                f"+{additions} / -{deletions}",
                entry.get("area", ""),
                ", ".join(entry.get("related_items", []) or []),
            ]
        )
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_commits(ws, commits):
    headers = ["해시", "날짜", "작성자", "제목"]
    ws.append(headers)
    style_header(ws, [14, 20, 20, 90])
    for commit in commits:
        ws.append(
            [
                commit.get("hash", ""),
                commit.get("date", ""),
                commit.get("author", ""),
                commit.get("subject", ""),
            ]
        )
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_logs(ws, logs):
    headers = ["명령", "종료 코드", "상태", "비고"]
    ws.append(headers)
    style_header(ws, [80, 10, 12, 40])
    for log in logs:
        ws.append(
            [
                log.get("command", ""),
                log.get("exit_code", ""),
                log.get("status", ""),
                log.get("note", ""),
            ]
        )
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main():
    args = parse_args()
    data = load_input(args.input)
    items = normalize_items(data["items"])
    generated_at = data.get("generated_at") or datetime.now().astimezone().isoformat(timespec="seconds")
    date_part = args.date or generated_at[:10]
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    stem = args.name or f"release-check-{date_part}"
    output = outdir / f"{stem}.xlsx"

    workbook = Workbook()
    summary = workbook.active
    summary.title = "요약"
    write_summary(summary, data, items, generated_at)
    write_items(workbook.create_sheet("점검 항목"), items)
    write_files(workbook.create_sheet("변경 파일"), data.get("files", []))
    write_commits(workbook.create_sheet("커밋"), data.get("commits", []))
    write_logs(workbook.create_sheet("실행 로그"), data.get("logs", []))
    workbook.save(output)

    print(
        json.dumps(
            {
                "xlsx": str(output),
                "items": len(items),
                "priority": dict(Counter(i["priority"] for i in items)),
                "audience": dict(Counter(i["audience"] for i in items)),
                "result": dict(Counter(i["status"] for i in items)),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
